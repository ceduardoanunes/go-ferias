#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Atualiza a ADMISSAO dos colaboradores no backend (Render/local) a partir das
abas POR SETOR da planilha `docs/colaboradores_por_departamento.xlsx`.

Por que as abas por setor (e nao a aba consolidada "Colaboradores"): a aba
consolidada esta velha (igual ao pull de 20/08); as abas de setor foram as que
o RH reconciliou (datas de admissao corrigidas p/ a data historica, desfazendo
o efeito da troca de CNPJ). Casa por `ID (sistema)` -> so faz PATCH em quem
realmente difere do estado VIVO da API. NAO cria ninguem e nao toca em periodos.

Uso (dry-run, so mostra o que faria):
  py atualizar_admissao.py --url https://goferias.onrender.com --senha demo

Para gravar de verdade, acrescente --apply:
  py atualizar_admissao.py --url https://goferias.onrender.com --senha demo --apply

Opcoes:
  --xlsx      caminho da planilha (padrao: docs/colaboradores_por_departamento.xlsx)
  --url       base do backend (obrigatorio). Ex.: https://goferias.onrender.com
  --email     login admin (padrao: admin@goegrow.com.br)
  --senha     senha admin (padrao: demo)
  --apply     grava (PATCH). Sem isso, so simula (dry-run).
  --inseguro  ignora verificacao do certificado TLS (SSL do Python no Windows).

Seguro re-rodar: idempotente (so mexe em quem difere). Linhas SEM ID (ex.: Sara
Efigenia, que nao existe no sistema) sao apenas RELATADAS p/ lancar a mao.
"""
import os, json, argparse, ssl, urllib.request, urllib.error

try:
    import openpyxl
except ImportError:
    raise SystemExit("Falta openpyxl. Instale com:  py -m pip install openpyxl")

ABA_CONSOLIDADA = "Colaboradores"   # aba resumo (dados velhos) -> ignorada
COL_ID = "ID (sistema)"
COL_ADM = "Admissão"
COL_NOME = "Nome"


def dia(v):
    """Normaliza a celula de data para 'YYYY-MM-DD' (ou None)."""
    if v is None:
        return None
    if hasattr(v, "date"):
        return v.date().isoformat()
    s = str(v).strip()
    return s or None


def http(url, token=None, method="GET", body=None, ctx=None):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, context=ctx) as r:
            txt = r.read().decode("utf-8")
            return (r.status, json.loads(txt) if txt else None)
    except urllib.error.HTTPError as e:
        txt = e.read().decode("utf-8", "ignore")
        try:
            msg = json.loads(txt).get("message", txt)
        except Exception:
            msg = txt
        raise SystemExit(f"Erro HTTP {e.code} em {method} {url}: {msg}")
    except urllib.error.URLError as e:
        raise SystemExit(f"Nao consegui conectar em {url}: {e.reason}")


def ler_planilha(caminho):
    """Le as abas POR SETOR -> {id: {'nome','adm','aba'}} e lista de linhas sem id."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    por_id, sem_id = {}, []
    for aba in wb.sheetnames:
        if aba == ABA_CONSOLIDADA:
            continue
        ws = wb[aba]
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            continue
        hdr = list(linhas[0])
        try:
            iid, iadm, inome = hdr.index(COL_ID), hdr.index(COL_ADM), hdr.index(COL_NOME)
        except ValueError:
            print(f"  ! aba '{aba}' sem cabecalho esperado -> pulada")
            continue
        for r in linhas[1:]:
            if not any(x is not None for x in r):
                continue
            rid = r[iid]
            nome = (r[inome] or "").strip()
            adm = dia(r[iadm])
            if not rid:
                sem_id.append({"nome": nome, "adm": adm, "aba": aba})
                continue
            if rid in por_id and por_id[rid]["adm"] != adm:
                print(f"  ! ID {rid} aparece em 2 abas com admissao diferente "
                      f"({por_id[rid]['adm']} x {adm}) - usando '{aba}'")
            por_id[rid] = {"nome": nome, "adm": adm, "aba": aba}
    return por_id, sem_id


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=os.path.join("..", "..", "docs",
                    "colaboradores_por_departamento.xlsx"),
                    help="planilha (padrao: docs/colaboradores_por_departamento.xlsx)")
    ap.add_argument("--url", required=True, help="base do backend, ex.: https://goferias.onrender.com")
    ap.add_argument("--email", default="admin@goegrow.com.br")
    ap.add_argument("--senha", default="demo")
    ap.add_argument("--apply", action="store_true", help="grava (PATCH). Sem isso: dry-run.")
    ap.add_argument("--inseguro", action="store_true", help="ignora verificacao do certificado TLS")
    a = ap.parse_args()

    if not os.path.exists(a.xlsx):
        raise SystemExit(f"Planilha nao encontrada: {a.xlsx}")

    base = a.url.rstrip("/")
    ctx = ssl._create_unverified_context() if a.inseguro else None

    por_id, sem_id = ler_planilha(a.xlsx)
    print(f"Planilha: {len(por_id)} colaborador(es) com ID nas abas de setor; "
          f"{len(sem_id)} sem ID.\n")

    # login + estado VIVO
    _, login = http(f"{base}/auth/login", method="POST",
                    body={"email": a.email, "senha": a.senha}, ctx=ctx)
    token = login["token"]
    print(f"Login OK como {login['nome']} ({login['papel']}) em {base}.")
    _, atuais = http(f"{base}/colaboradores", token=token, ctx=ctx)
    vivo = {c["id"]: c for c in (atuais or [])}
    print(f"{len(vivo)} colaborador(es) no sistema.\n")

    mudancas, iguais, ausentes = [], 0, []
    for rid, info in por_id.items():
        alvo = info["adm"]
        c = vivo.get(rid)
        if c is None:
            ausentes.append(info)
            continue
        if not alvo:
            continue                     # planilha sem data -> nao mexe
        if (c.get("admissao") or None) == alvo:
            iguais += 1
            continue
        mudancas.append((rid, info["nome"], c.get("admissao"), alvo))

    modo = "APLICANDO" if a.apply else "DRY-RUN (nada sera gravado)"
    print(f"=== {modo} ===")
    print(f"{len(mudancas)} admissao(oes) a mudar | {iguais} ja corretas | "
          f"{len(ausentes)} ID(s) da planilha ausente(s) no sistema | "
          f"{len(sem_id)} linha(s) sem ID\n")

    for rid, nome, de, para in sorted(mudancas, key=lambda x: x[1]):
        print(f"  {nome:42s} {str(de):>10s} -> {para}")
        if a.apply:
            http(f"{base}/colaboradores/{rid}", token=token, method="PATCH",
                 body={"admissao": para}, ctx=ctx)

    if ausentes:
        print("\nIDs da planilha que NAO existem no sistema (ignorados):")
        for i in ausentes:
            print(f"  - {i['nome']} (aba {i['aba']}, adm {i['adm']})")
    if sem_id:
        print("\nLinhas SEM ID (criar a mao no app - nao sao tocadas aqui):")
        for i in sem_id:
            print(f"  - {i['nome']} (aba {i['aba']}, adm {i['adm']})")

    if not a.apply and mudancas:
        print("\n>>> Isto foi um DRY-RUN. Para gravar, rode de novo com --apply")
    print("\nOk.")


if __name__ == "__main__":
    main()
